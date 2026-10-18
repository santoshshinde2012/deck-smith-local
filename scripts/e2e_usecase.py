"""The DeckSmith use case, end to end against a running API, with a pass/fail table.

An analyst (bob) asks for a churn review deck, edits and then approves the outline, and downloads an
editable PPTX whose every number traces back to a fact ID. alice (EU-West only) gets a deck with her
own numbers, carol (no data access) is refused, and nobody can read someone else's run.

Usage: python scripts/e2e_usecase.py [--api http://localhost:8000] [--out out]
Exits non-zero if any step fails; writes out/e2e-usecase.md and out/usecase-<user>.pptx.
"""

from __future__ import annotations

import argparse
import math
import re
import sys
import time
from collections.abc import Callable
from pathlib import Path
from typing import Any

import httpx
from pptx import Presentation

from decksmith.adapters.local.design.charts import category_labels
from decksmith.domain.deckspec import DeckSpec
from decksmith.domain.facts import FactRegistry, format_value
from decksmith.domain.integrity import FACT_PATTERN, chart_consistency, numeric_integrity

REQUEST = "Build the Q3 2026 customer retention review for the exec team"
NUMBER = re.compile(r"(?<![\w.])[$€£]?\d[\d,]*(?:\.\d+)?%?")
Row = tuple[str, bool, str]


class UseCase:
    def __init__(self, api: str, out: Path) -> None:
        self.api, self.out, self.rows, self.ctx = api.rstrip("/"), out, list[Row](), dict[str, Any]()

    def call(self, user: str | None, method: str, path: str, **kw: Any) -> httpx.Response:
        headers = {"Authorization": f"Bearer {user}-token"} if user else {}
        return httpx.request(method, self.api + path, headers=headers, timeout=180, **kw)

    def step(self, name: str, fn: Callable[[], str]) -> None:
        started = time.monotonic()
        try:
            detail, ok = fn(), True
        except Exception as exc:  # each step reports its own failure and the run continues
            detail, ok = f"{type(exc).__name__}: {exc}", False
        self.rows.append((name, ok, f"{detail} ({time.monotonic() - started:.1f}s)"))

    # --- steps -------------------------------------------------------------------------------
    def health(self) -> str:
        r = httpx.get(self.api + "/healthz", timeout=10)
        assert r.status_code == 200, r.status_code
        return f"API up at {self.api}"

    def request_outline(self, user: str) -> str:
        view = self.call(user, "POST", "/v1/decks", json={"request": REQUEST}).json()
        assert view.get("status") == "awaiting_approval", view
        slides = [s["slide_id"] for s in view["outline"]["slides"]]
        assert len(slides) >= 3, slides
        self.ctx[f"{user}_run"] = view["run_id"]
        self.ctx[f"{user}_outline"] = view["outline"]
        return f"run {view['run_id'][:8]}: outline {', '.join(slides)}, paused for approval"

    def edit_outline(self) -> str:
        run = self.ctx["bob_run"]
        view = self.call(
            "bob", "POST", f"/v1/decks/{run}/decision", json={"action": "edit", "feedback": "drop the segment slide"}
        ).json()
        slides = [s["slide_id"] for s in view["outline"]["slides"]]
        assert view["status"] == "awaiting_approval" and "s04" not in slides, view
        assert view["run_id"] == run and not view.get("facts") and not view.get("artifact"), view
        self.ctx["bob_outline"] = view["outline"]
        return f"edit applied ({', '.join(slides)}); still waiting for approval"

    def approve(self, user: str) -> str:
        run = self.ctx[f"{user}_run"]
        view = self.call(user, "POST", f"/v1/decks/{run}/decision", json={"action": "approve"}).json()
        assert view["status"] == "done", view
        assert view["run_id"] == run, "approval must resume the edited run"
        expected = [(s["slide_id"], s["layout"]) for s in self.ctx[f"{user}_outline"]["slides"]]
        actual = [(s["slide_id"], s["layout"]) for s in view["spec"]["slides"]]
        assert actual == expected, (actual, expected)
        if user == "bob":
            assert len(actual) == 4 and all(f["source"]["template_id"] != "churn_by_segment" for f in view["facts"])
        self.ctx[f"{user}_view"] = view
        return f"status done: {view['artifact']['slide_count']} slides, {len(view['facts'])} facts"

    def download(self, user: str) -> str:
        run = self.ctx[f"{user}_run"]
        r = self.call(user, "GET", f"/v1/decks/{run}/file")
        assert r.status_code == 200 and r.content[:2] == b"PK", r.status_code
        path = self.out / f"usecase-{user}.pptx"
        path.write_bytes(r.content)
        deck = Presentation(str(path))
        charts = sum(1 for s in deck.slides for sh in s.shapes if getattr(sh, "has_chart", False))
        assert len(deck.slides) == self.ctx[f"{user}_view"]["artifact"]["slide_count"]
        assert charts >= 1, "expected native, editable charts"
        self.ctx[f"{user}_deck"] = deck
        chart_label = "chart" if charts == 1 else "charts"
        return f"{path} ({len(r.content):,} bytes, {len(deck.slides)} slides, {charts} native {chart_label})"

    def numbers_trace_to_facts(self, user: str) -> str:
        view, deck = self.ctx[f"{user}_view"], self.ctx[f"{user}_deck"]
        spec, registry = DeckSpec.model_validate(view["spec"]), FactRegistry.load(view["facts"])
        assert numeric_integrity(spec, registry) == [] and chart_consistency(spec, registry) == []
        allowed = {format_value(f.value, f.format, f.unit) for f in registry.all() if f.value is not None}
        seen, sources_text, chart_points = [], "", 0
        assert len(deck.slides) == len(spec.slides)
        for slide, slide_spec in zip(deck.slides, spec.slides, strict=True):
            text = "\n".join(sh.text_frame.text for sh in slide.shapes if sh.has_text_frame)
            if slide_spec.layout == "appendix_sources":
                sources_text += text
                continue
            for number in NUMBER.findall(text):
                assert number in allowed, f"{number!r} is not a formatted scalar fact"
                seen.append(number)
            charts = [shape.chart for shape in slide.shapes if shape.has_chart]
            assert len(charts) == (1 if slide_spec.chart else 0)
            if slide_spec.chart:
                chart = charts[0]
                expected_categories = list(registry.get(slide_spec.chart.series[0].fact_id).categories)
                assert [category.label for category in chart.plots[0].categories] == category_labels(
                    expected_categories
                )
                assert len(chart.series) == len(slide_spec.chart.series)
                for actual_series, series_spec in zip(chart.series, slide_spec.chart.series, strict=True):
                    fact = registry.get(series_spec.fact_id)
                    expected_name = FACT_PATTERN.sub(
                        lambda match: format_value(
                            registry.get(match.group(1)).value,  # type: ignore[arg-type]
                            registry.get(match.group(1)).format,
                            registry.get(match.group(1)).unit,
                        ),
                        series_spec.name,
                    )
                    assert actual_series.name == expected_name
                    assert list(fact.categories) == expected_categories
                    assert len(actual_series.values) == len(fact.values)
                    assert all(
                        math.isclose(actual, expected, rel_tol=1e-12, abs_tol=1e-12)
                        for actual, expected in zip(actual_series.values, fact.values, strict=True)
                    ), f"chart values differ from {fact.id}"
                    chart_points += len(fact.values)
            notes = slide.notes_slide.notes_text_frame.text
            for fact_id in slide_spec.fact_refs():
                assert fact_id in notes, f"notes miss provenance for {fact_id}"
        missing = [f.id for f in registry.all() if f"{f.id} " not in sources_text + " "]
        assert sources_text and not missing, f"sources slide misses {missing}"
        return (
            f"{len(seen)} scalar numbers and {chart_points} chart points match facts; "
            f"sources slide cites {len(registry)} fact IDs"
        )

    def alice_deck(self) -> str:
        run = self.request_outline("alice").split(":")[0]
        return f"{run}: {self.approve('alice')}"

    def alice_consistent(self) -> str:
        path = self.download("alice").split(" (")[0]
        return f"{self.numbers_trace_to_facts('alice')}; saved {path}"

    def permissions_differ(self) -> str:
        def active(user: str) -> float:
            facts = self.ctx[f"{user}_view"]["facts"]
            return float(next(f["value"] for f in facts if f["label"] == "Active subscriptions"))

        bob, alice = active("bob"), active("alice")
        assert 0 < alice < bob, (alice, bob)
        return f"active subscriptions: bob {bob:,.0f} (all regions) vs alice {alice:,.0f} (EU-West only)"

    def carol_refused(self) -> str:
        self.request_outline("carol")
        run = self.ctx["carol_run"]
        view = self.call("carol", "POST", f"/v1/decks/{run}/decision", json={"action": "approve"}).json()
        assert view["status"] == "no_data" and not view.get("artifact"), view
        return f"no_data: {view['error']}"

    def isolation(self) -> str:
        run = self.ctx["bob_run"]
        codes = (
            self.call("alice", "GET", f"/v1/decks/{run}").status_code,
            self.call("alice", "GET", f"/v1/decks/{run}/file").status_code,
            self.call(None, "POST", "/v1/decks", json={"request": REQUEST}).status_code,
        )
        assert codes == (404, 404, 401), codes
        return "alice cannot see bob's run or file (404); anonymous request rejected (401)"

    def run(self) -> int:
        self.out.mkdir(parents=True, exist_ok=True)
        steps: list[tuple[str, Callable[[], str]]] = [
            ("1. API is up", self.health),
            ("2. bob asks for a churn review deck", lambda: self.request_outline("bob")),
            ("3. bob edits the outline (human in the loop)", self.edit_outline),
            ("4. bob approves; data is fetched as bob", lambda: self.approve("bob")),
            ("5. bob downloads an editable PPTX", lambda: self.download("bob")),
            ("6. scalar and chart values match facts", lambda: self.numbers_trace_to_facts("bob")),
            ("7. alice (EU-West only) gets her own deck", self.alice_deck),
            ("8. alice's deck is consistent too", self.alice_consistent),
            ("9. permissions change the numbers", self.permissions_differ),
            ("10. carol (no access) is refused", self.carol_refused),
            ("11. runs are private; auth is required", self.isolation),
        ]
        for name, fn in steps:
            self.step(name, fn)
        width = max(len(n) for n, _, _ in self.rows)
        lines = ["| Step | Result | Detail |", "| --- | --- | --- |"]
        print(f"\n{'Step':<{width}}  Result  Detail")
        for name, ok, detail in self.rows:
            print(f"{name:<{width}}  {'PASS' if ok else 'FAIL':<6}  {detail}")
            lines.append(f"| {name} | {'PASS' if ok else 'FAIL'} | {detail.replace('|', '/')} |")
        passed = sum(ok for _, ok, _ in self.rows)
        summary = f"{passed}/{len(self.rows)} steps passed"
        print(f"\n{summary}")
        report = "# DeckSmith use case, end to end\n\n" + summary + "\n\n" + "\n".join(lines) + "\n"
        (self.out / "e2e-usecase.md").write_text(report)
        return 0 if passed == len(self.rows) else 1


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--api", default="http://localhost:8000")
    parser.add_argument("--out", type=Path, default=Path("out"))
    args = parser.parse_args()
    sys.exit(UseCase(args.api, args.out).run())


if __name__ == "__main__":
    main()
