"""Build a contact sheet from a DeckSmith PPTX.

Usage: python scripts/preview_deck.py docs/examples/sample-deck.pptx OUTPUT.png

Requires Pillow and matplotlib. This is a code-rendered preview, not a pixel-exact
PowerPoint screenshot: the host may substitute the deck's Georgia and Calibri faces.
Text, shape positions, colours, chart categories and series values come from the PPTX.
The PNG metadata records the source PPTX hash and rendering method.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import math
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
from matplotlib import pyplot as plt
from PIL import Image, ImageDraw, ImageFont, PngImagePlugin
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

EMU_PER_INCH = 914400
SLIDE_WIDTH = 960
GAP = 28
CANVAS = "#E7E2DA"
INK = "#1C1915"


def font_file(name: str | None, bold: bool) -> str:
    serif = (bold and "NotoSerif-Bold.ttf") or "NotoSerif-Regular.ttf"
    sans = (bold and "Inter-Bold.ttf") or "Inter-Regular.ttf"
    chosen = serif if name and "georgia" in name.lower() else sans
    roots = [Path("/usr/share/fonts/truetype/noto"), Path("/usr/share/fonts/truetype/macos")]
    for root in roots:
        candidate = root / chosen
        if candidate.exists():
            return str(candidate)
    fallback = "DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf"
    return str(Path("/usr/share/fonts/truetype/dejavu") / fallback)


def colour_of(colour: object, default: str = INK) -> str:
    try:
        return f"#{colour.rgb}"  # type: ignore[attr-defined]
    except (AttributeError, TypeError):
        return default


def fill_hex(fill: object, default: str | None = None) -> str | None:
    try:
        if fill.type is None:  # type: ignore[attr-defined]
            return default
        return colour_of(fill.fore_color, default or INK)  # type: ignore[attr-defined]
    except (AttributeError, TypeError):
        return default


def wrap(text: str, font: ImageFont.FreeTypeFont, width: float) -> list[str]:
    if not text:
        return []
    lines: list[str] = []
    for paragraph in text.split("\n"):
        current = ""
        for word in paragraph.split() or [""]:
            proposal = word if not current else f"{current} {word}"
            if current and font.getlength(proposal) > width:
                lines.append(current)
                current = word
            else:
                current = proposal
        lines.append(current)
    return lines


def slide_background(slide: object) -> str:
    try:
        found = fill_hex(slide.background.fill)  # type: ignore[attr-defined]
    except (AttributeError, TypeError):
        found = None
    return found or "#FFFFFF"


def draw_shape(image: Image.Image, shape: object, scale: float) -> None:
    kind = shape.shape_type  # type: ignore[attr-defined]
    if kind not in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.FREEFORM):
        return
    left, top = round(shape.left * scale), round(shape.top * scale)  # type: ignore[attr-defined]
    width, height = max(1, round(shape.width * scale)), max(1, round(shape.height * scale))  # type: ignore[attr-defined]
    fill = fill_hex(shape.fill)  # type: ignore[attr-defined]
    line = None
    try:
        line = colour_of(shape.line.color)  # type: ignore[attr-defined]
        line_width = max(1, round(shape.line.width.pt * scale * EMU_PER_INCH / 72))  # type: ignore[attr-defined]
    except (AttributeError, TypeError):
        line_width = 1
    draw = ImageDraw.Draw(image)
    autoshape = getattr(shape, "auto_shape_type", None)
    if autoshape == MSO_SHAPE.OVAL:
        if fill:
            draw.ellipse((left, top, left + width, top + height), fill=fill, outline=line, width=line_width)
        elif line:
            draw.ellipse((left, top, left + width, top + height), outline=line, width=line_width)
        return
    radius = 0
    if autoshape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            radius = max(2, round(min(width, height) * float(shape.adjustments[0]) * 0.5))  # type: ignore[attr-defined]
        except (AttributeError, IndexError, TypeError, ValueError):
            radius = max(2, round(min(width, height) * 0.08))
    box = (left, top, left + width, top + height)
    if fill:
        draw.rounded_rectangle(box, radius=radius, fill=fill, outline=line, width=line_width if line else 1)
    elif line:
        draw.rounded_rectangle(box, radius=radius, outline=line, width=line_width)


def draw_text(image: Image.Image, shape: object, scale: float) -> None:
    frame = shape.text_frame  # type: ignore[attr-defined]
    if not any(paragraph.text for paragraph in frame.paragraphs):
        return
    points_to_px = scale * EMU_PER_INCH / 72
    x = (shape.left + frame.margin_left) * scale  # type: ignore[attr-defined]
    y = (shape.top + frame.margin_top) * scale  # type: ignore[attr-defined]
    width = max(1.0, (shape.width - frame.margin_left - frame.margin_right) * scale)  # type: ignore[attr-defined]
    height = max(1.0, (shape.height - frame.margin_top - frame.margin_bottom) * scale)  # type: ignore[attr-defined]
    paragraphs = []
    for paragraph in frame.paragraphs:
        run = paragraph.runs[0] if paragraph.runs else None
        size = run.font.size if run is not None else paragraph.font.size
        size_pt = size.pt if size is not None else 16
        bold = bool(run and run.font.bold)
        name = run.font.name if run is not None else None
        font = ImageFont.truetype(font_file(name, bold), max(1, round(size_pt * points_to_px)))
        lines = wrap(paragraph.text, font, width) if frame.word_wrap is not False else [paragraph.text]
        centered = paragraph.alignment == PP_ALIGN.CENTER or paragraph.alignment == PP_ALIGN.RIGHT
        align = paragraph.alignment
        colour = colour_of(run.font.color, INK) if run is not None else INK
        paragraphs.append((lines, font, align, centered, colour, round(size_pt * points_to_px * 1.05)))
    text_height = sum(len(lines) * step for lines, _, _, _, _, step in paragraphs)
    if frame.vertical_anchor == MSO_ANCHOR.MIDDLE:
        y += max(0, (height - text_height) / 2)
    elif frame.vertical_anchor == MSO_ANCHOR.BOTTOM:
        y += max(0, height - text_height)
    draw = ImageDraw.Draw(image)
    for lines, font, align, _, colour, step in paragraphs:
        for line in lines:
            if align == PP_ALIGN.CENTER:
                left = x + max(0, (width - font.getlength(line)) / 2)
            elif align == PP_ALIGN.RIGHT:
                left = x + max(0, width - font.getlength(line))
            else:
                left = x
            draw.text((left, y), line, font=font, fill=colour, anchor="lt")
            y += step


def series_colour(series: object, index: int, fallback: str) -> str:
    try:
        point = series.points[index]  # type: ignore[attr-defined]
        for candidate in (point.marker.format.fill, point.format.fill):
            found = fill_hex(candidate)
            if found:
                return found
    except (AttributeError, IndexError, TypeError):
        pass
    try:
        return colour_of(series.format.line.color, fallback)  # type: ignore[attr-defined]
    except (AttributeError, TypeError):
        return fallback


def draw_chart(image: Image.Image, shape: object, scale: float) -> None:
    chart = shape.chart  # type: ignore[attr-defined]
    width, height = max(1, round(shape.width * scale)), max(1, round(shape.height * scale))  # type: ignore[attr-defined]
    figure, axes = plt.subplots(figsize=(width / 100, height / 100), dpi=100)
    figure.patch.set_facecolor("#FBF9F5")
    axes.set_facecolor("#FBF9F5")
    categories = [str(category) for category in chart.plots[0].categories]
    palette = ["#1A3144", "#9C4221", "#3E6A74", "#7A5E24"]
    percent = False
    for series_index, series in enumerate(chart.series):
        values = [float(value) for value in series.values]
        percent = percent or (max(values) <= 1.5 and any(value < 1 for value in values))
        colours = [series_colour(series, index, palette[series_index % len(palette)]) for index in range(len(values))]
        if chart.chart_type == XL_CHART_TYPE.LINE_MARKERS:
            axes.plot(categories, values, color=colours[0], linewidth=2.0, solid_capstyle="round")
            axes.scatter(categories, values, c=colours, s=28, zorder=3)
            if not chart.value_axis.visible:
                for category, value, colour in zip(categories, values, colours, strict=True):
                    axes.annotate(
                        f"{value * 100:.1f}%" if percent else f"{value:,.0f}",
                        (category, value),
                        textcoords="offset points",
                        xytext=(0, 8),
                        ha="center",
                        fontsize=8,
                        color=colour,
                    )
        elif chart.chart_type in (XL_CHART_TYPE.BAR_CLUSTERED, XL_CHART_TYPE.COLUMN_CLUSTERED):
            positions = list(range(len(categories)))
            if chart.chart_type == XL_CHART_TYPE.BAR_CLUSTERED:
                axes.barh(positions, values, color=colours, height=0.62)
                axes.set_yticks(positions, categories)
                axes.invert_yaxis()
            else:
                axes.bar(positions, values, color=colours, width=0.62)
                axes.set_xticks(positions, categories)
            if not chart.value_axis.visible:
                for position, value, colour in zip(positions, values, colours, strict=True):
                    label = f"{value * 100:.1f}%" if percent else f"{value:,.0f}"
                    if chart.chart_type == XL_CHART_TYPE.BAR_CLUSTERED:
                        axes.annotate(
                            label,
                            (value, position),
                            textcoords="offset points",
                            xytext=(6, 0),
                            va="center",
                            fontsize=8,
                            color=colour,
                        )
                    else:
                        axes.annotate(
                            label,
                            (position, value),
                            textcoords="offset points",
                            xytext=(0, 6),
                            ha="center",
                            fontsize=8,
                            color=colour,
                        )
        else:
            raise ValueError(f"Unsupported chart in sample preview: {chart.chart_type}")
    axes.tick_params(labelsize=9, length=0, colors="#1C1915")
    if not chart.value_axis.visible:
        if chart.chart_type == XL_CHART_TYPE.BAR_CLUSTERED:
            peak = max(values) if values else 1
            floor = min(values) if values else 0
            axes.set_xlim(min(0, floor), peak * 1.35 if peak > 0 else 1)
            axes.xaxis.set_visible(False)
            axes.spines[["top", "right", "bottom"]].set_visible(False)
        else:
            axes.yaxis.set_visible(False)
            axes.spines[["top", "right", "left"]].set_visible(False)
    else:
        axes.grid(axis="y", color="#E3DCD0", linewidth=0.6)
        axes.spines[["top", "right"]].set_visible(False)
    for spine in axes.spines.values():
        spine.set_color("#E3DCD0")
    figure.subplots_adjust(left=0.16, bottom=0.16, right=0.96, top=0.92)
    if chart.has_legend:
        axes.legend([series.name for series in chart.series], frameon=False, fontsize=8)
    buffer = io.BytesIO()
    figure.savefig(buffer, format="png", dpi=100)
    plt.close(figure)
    buffer.seek(0)
    with Image.open(buffer) as rendered:
        image.paste(rendered.convert("RGB"), (round(shape.left * scale), round(shape.top * scale)))


def render(source: Path, destination: Path) -> None:
    presentation = Presentation(source)
    scale = SLIDE_WIDTH / presentation.slide_width
    slide_height = round(presentation.slide_height * scale)
    columns = 3 if len(presentation.slides) > 1 else 1
    rows = math.ceil(len(presentation.slides) / columns)
    width = SLIDE_WIDTH * columns + GAP * (columns + 1)
    height = slide_height * rows + GAP * (rows + 1)
    contact = Image.new("RGB", (width, height), CANVAS)
    for index, slide in enumerate(presentation.slides):
        image = Image.new("RGB", (SLIDE_WIDTH, slide_height), slide_background(slide))
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.CHART:
                draw_chart(image, shape, scale)
            else:
                draw_shape(image, shape, scale)
                if shape.has_text_frame:
                    draw_text(image, shape, scale)
        origin = (
            GAP + (index % columns) * (SLIDE_WIDTH + GAP),
            GAP + (index // columns) * (slide_height + GAP),
        )
        contact.paste(image, origin)
    metadata = PngImagePlugin.PngInfo()
    metadata.add_text(
        "Description", "Code-rendered from actual PPTX shapes, text and series; not a native Office screenshot"
    )
    metadata.add_text("Source-SHA256", hashlib.sha256(source.read_bytes()).hexdigest())
    metadata.add_text("Slide-Count", str(len(presentation.slides)))
    destination.parent.mkdir(parents=True, exist_ok=True)
    contact.save(destination, pnginfo=metadata)
    print(f"{destination}: {contact.width} x {contact.height}; {len(presentation.slides)} slides")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("destination", type=Path)
    arguments = parser.parse_args()
    render(arguments.source, arguments.destination)
