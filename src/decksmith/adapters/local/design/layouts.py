"""One painter per slide layout. A new layout is a new painter registered here, not a branch in the renderer."""

from __future__ import annotations

from typing import Any, Protocol

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from decksmith.adapters.local.design.canvas import (
    Canvas,
    add_shape,
    add_text,
    card_spans,
    folio,
    footer_sources,
    paint_action_title,
    paint_background,
    paint_footer,
)
from decksmith.adapters.local.design.charts import add_chart, category_labels, direction_word
from decksmith.adapters.local.design.fitting import fit_paragraphs, single_line, too_dense
from decksmith.adapters.local.design.motifs import paint_cover
from decksmith.adapters.local.design.text import fill_text
from decksmith.adapters.local.design.theme import Theme
from decksmith.domain.deckspec import Kpi, Slide
from decksmith.domain.errors import GuardrailViolation
from decksmith.domain.facts import Fact, FactRegistry, format_value


class SlidePainter(Protocol):
    """Renders one layout. Returns whether body text still overflowed at the minimum size."""

    layout: str

    def paint(
        self, slide: Any, spec: Slide, facts: FactRegistry, theme: Theme, canvas: Canvas, number: int
    ) -> bool: ...


class LayoutRegistry:
    """Open/closed collection of painters, keyed by the DeckSpec layout name."""

    def __init__(self, painters: tuple[SlidePainter, ...] | list[SlidePainter]) -> None:
        self._painters: dict[str, SlidePainter] = {}
        for painter in painters:
            self._painters[painter.layout] = painter

    def layouts(self) -> tuple[str, ...]:
        return tuple(sorted(self._painters))

    def paint(self, slide: Any, spec: Slide, facts: FactRegistry, theme: Theme, canvas: Canvas, number: int) -> bool:
        return self._painters[spec.layout].paint(slide, spec, facts, theme, canvas, number)


class TitleLayout:
    layout = "title"

    def paint(self, slide: Any, spec: Slide, facts: FactRegistry, theme: Theme, canvas: Canvas, number: int) -> bool:
        return paint_cover(slide, spec, facts, theme, canvas, number)


class SectionLayout:
    layout = "section"

    def paint(self, slide: Any, spec: Slide, facts: FactRegistry, theme: Theme, canvas: Canvas, number: int) -> bool:
        paint_background(slide, theme.cover)
        left, width = canvas.span(0, 9)
        add_shape(slide, MSO_SHAPE.RECTANGLE, left, int(canvas.height * 0.36), 365760, 27432, theme.accent)
        title = fill_text(spec.title, facts)
        fit = fit_paragraphs([title], canvas.inches(width), 1.6, max_pt=40, min_pt=24)
        add_text(
            slide,
            "\n".join(fit.lines),
            left,
            int(canvas.height * 0.40),
            width,
            1463040,
            font=theme.font_display,
            size=fit.font_pt,
            color=theme.cover_ink,
        )
        overflow = fit.overflow
        subtitle = fill_text(spec.subtitle, facts)
        if subtitle:
            sub = fit_paragraphs([subtitle], canvas.inches(width), 0.6, max_pt=18, min_pt=14)
            add_text(
                slide,
                "\n".join(sub.lines),
                left,
                int(canvas.height * 0.40) + 1554480,
                width,
                548640,
                font=theme.font_body,
                size=sub.font_pt,
                color=theme.cover_muted,
            )
            overflow = overflow or sub.overflow
        paint_footer(slide, canvas, theme, number, "", inverse=True)
        return overflow


class AgendaLayout:
    layout = "agenda"

    def paint(self, slide: Any, spec: Slide, facts: FactRegistry, theme: Theme, canvas: Canvas, number: int) -> bool:
        paint_background(slide, theme.paper)
        overflow = paint_action_title(slide, fill_text(spec.title, facts), theme, canvas)
        lines = [fill_text(bullet, facts) for bullet in spec.bullets]
        overflow = overflow or _rows(slide, lines, theme, canvas, numbered=True) or too_dense(lines)
        paint_footer(slide, canvas, theme, number, _sources(spec, facts))
        return overflow


class KpiLayout:
    layout = "kpi"

    def paint(self, slide: Any, spec: Slide, facts: FactRegistry, theme: Theme, canvas: Canvas, number: int) -> bool:
        paint_background(slide, theme.paper)
        overflow = paint_action_title(slide, fill_text(spec.title, facts), theme, canvas)
        spans = card_spans(len(spec.kpis))
        top = 1417320  # 1.55 in
        height = 2377440  # 2.60 in
        for (start, count), kpi in zip(spans, spec.kpis, strict=True):
            left, width = canvas.span(start, count)
            overflow = overflow or _kpi_card(slide, kpi, facts, theme, canvas, left, top, width, height)
        lines = [fill_text(bullet, facts) for bullet in spec.bullets]
        if lines:
            callout_top = top + height + 137160
            callout_height = min(1188720, canvas.footer_top - callout_top - 91440)
            if callout_height > 274320:
                overflow = overflow or _callout(slide, lines, theme, canvas, callout_top, callout_height)
            overflow = overflow or too_dense(lines)
        paint_footer(slide, canvas, theme, number, _sources(spec, facts))
        return overflow


class ChartLayout:
    layout = "chart"

    def paint(self, slide: Any, spec: Slide, facts: FactRegistry, theme: Theme, canvas: Canvas, number: int) -> bool:
        paint_background(slide, theme.paper)
        overflow = paint_action_title(slide, fill_text(spec.title, facts), theme, canvas)
        assert spec.chart is not None
        left, width = canvas.span(0, 12)
        top = 1554480
        height = canvas.footer_top - top - 91440
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height, theme.card, radius=0.04)
        categories, series, fmt = _series(spec, facts)
        pad = 137160
        add_chart(
            slide,
            spec.chart.type,
            categories,
            series,
            left + pad,
            top + pad,
            width - 2 * pad,
            height - 2 * pad,
            theme,
            fmt,
        )
        paint_footer(slide, canvas, theme, number, _sources(spec, facts))
        return overflow


class BulletsLayout:
    layout = "bullets"

    def paint(self, slide: Any, spec: Slide, facts: FactRegistry, theme: Theme, canvas: Canvas, number: int) -> bool:
        return _text_slide(slide, spec, facts, theme, canvas, number, columns=1)


class TwoContentLayout:
    layout = "two_content"

    def paint(self, slide: Any, spec: Slide, facts: FactRegistry, theme: Theme, canvas: Canvas, number: int) -> bool:
        return _text_slide(slide, spec, facts, theme, canvas, number, columns=2)


class ComparisonLayout:
    layout = "comparison"

    def paint(self, slide: Any, spec: Slide, facts: FactRegistry, theme: Theme, canvas: Canvas, number: int) -> bool:
        paint_background(slide, theme.paper)
        overflow = paint_action_title(slide, fill_text(spec.title, facts), theme, canvas)
        top = 1463040
        height = canvas.footer_top - top - 91440
        for index, column in enumerate(spec.columns):
            left, width = canvas.span(0 if index == 0 else 6, 6)
            lines = [fill_text(bullet, facts) for bullet in column.bullets]
            overflow = overflow or _filled_panel(
                slide,
                fill_text(column.heading, facts),
                lines,
                theme,
                canvas,
                left,
                top,
                width,
                height,
                folio(index + 1),
            )
        paint_footer(slide, canvas, theme, number, _sources(spec, facts))
        return overflow


class TimelineLayout:
    layout = "timeline"

    def paint(self, slide: Any, spec: Slide, facts: FactRegistry, theme: Theme, canvas: Canvas, number: int) -> bool:
        paint_background(slide, theme.paper)
        overflow = paint_action_title(slide, fill_text(spec.title, facts), theme, canvas)
        count = len(spec.milestones)
        left, total = canvas.span(0, 12)
        top = 1554480
        height = canvas.footer_top - top - 91440
        gap = canvas.gutter
        card_w = (total - gap * (count - 1)) // count
        rail = top + 274320
        add_shape(slide, MSO_SHAPE.RECTANGLE, left, rail, total, 27432, theme.rule)
        for index, milestone in enumerate(spec.milestones):
            card_left = left + index * (card_w + gap)
            add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, card_left, top, card_w, height, theme.card, radius=0.06)
            add_shape(slide, MSO_SHAPE.OVAL, card_left + 182880, rail - 45720, 137160, 137160, theme.accent)
            add_text(
                slide,
                folio(index + 1),
                card_left + 365760,
                top + 137160,
                card_w - 502920,
                411480,
                font=theme.font_display,
                size=22,
                color=theme.accent,
                bold=True,
            )
            label = fill_text(milestone.label, facts)
            detail = fill_text(milestone.detail, facts)
            label_fit = fit_paragraphs([label], canvas.inches(card_w - 320040), 0.8, max_pt=20, min_pt=14)
            detail_h = min(1371600, height // 3)
            detail_top = top + height - detail_h - 160020
            detail_fit = fit_paragraphs(
                [detail], canvas.inches(card_w - 320040), canvas.inches(detail_h - 91440), max_pt=16, min_pt=12
            )
            add_text(
                slide,
                "\n".join(label_fit.lines),
                card_left + 160020,
                top + 640080,
                card_w - 320040,
                731520,
                font=theme.font_display,
                size=label_fit.font_pt,
                color=theme.ink,
                bold=theme.display_bold,
            )
            add_shape(slide, MSO_SHAPE.RECTANGLE, card_left + 160020, detail_top - 137160, 365760, 27432, theme.accent)
            add_text(
                slide,
                "\n".join(detail_fit.lines),
                card_left + 160020,
                detail_top,
                card_w - 320040,
                detail_h,
                font=theme.font_body,
                size=detail_fit.font_pt,
                color=theme.muted,
                anchor=MSO_ANCHOR.TOP,
            )
            overflow = overflow or label_fit.overflow or detail_fit.overflow
        paint_footer(slide, canvas, theme, number, _sources(spec, facts))
        return overflow


class QuoteLayout:
    layout = "quote"

    def paint(self, slide: Any, spec: Slide, facts: FactRegistry, theme: Theme, canvas: Canvas, number: int) -> bool:
        paint_background(slide, theme.paper)
        overflow = paint_action_title(slide, fill_text(spec.title, facts), theme, canvas)
        left, width = canvas.span(0, 12)
        top = 1463040
        height = canvas.footer_top - top - 91440
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height, theme.card, radius=0.04)
        add_shape(slide, MSO_SHAPE.RECTANGLE, left, top, 91440, height, theme.accent)
        add_text(
            slide,
            "“",
            left + 274320,
            top + 91440,
            914400,
            822960,
            font=theme.font_display,
            size=60,
            color=theme.accent,
        )
        quote = fill_text(spec.quote, facts)
        fit = fit_paragraphs([quote], canvas.inches(width - 1371600), 2.2, max_pt=32, min_pt=18)
        add_text(
            slide,
            "\n".join(fit.lines),
            left + 320040,
            top + 914400,
            width - 640080,
            int(height * 0.5),
            font=theme.font_display,
            size=fit.font_pt,
            color=theme.ink,
        )
        overflow = overflow or fit.overflow
        attribution = fill_text(spec.attribution, facts)
        if attribution:
            add_shape(slide, MSO_SHAPE.RECTANGLE, left + 320040, top + height - 640080, 365760, 27432, theme.accent)
            add_text(
                slide,
                attribution,
                left + 777240,
                top + height - 731520,
                width - 1097280,
                365760,
                font=theme.font_body,
                size=16,
                color=theme.muted,
            )
        paint_footer(slide, canvas, theme, number, _sources(spec, facts))
        return overflow


class TableLayout:
    layout = "table"

    def paint(self, slide: Any, spec: Slide, facts: FactRegistry, theme: Theme, canvas: Canvas, number: int) -> bool:
        paint_background(slide, theme.paper)
        overflow = paint_action_title(slide, fill_text(spec.title, facts), theme, canvas)
        assert spec.table is not None
        columns = []
        categories: tuple[str, ...] = ()
        for column in spec.table.columns:
            fact = facts.get(column.fact_id)
            if fact.kind != "series":
                raise GuardrailViolation(f"Table needs a series fact: {fact.id}")
            categories = fact.categories
            columns.append((fill_text(column.heading, facts), fact))
        labels = category_labels(categories)
        left, width = canvas.span(0, 12)
        top = 1463040
        available = canvas.footer_top - top - 91440
        rows = max(1, len(labels))
        row_h = available // (rows + 1)
        overflow = overflow or row_h < 320040
        label_w = int(width * 0.28)
        value_w = (width - label_w) // max(1, len(columns))
        add_shape(slide, MSO_SHAPE.RECTANGLE, left, top, width, row_h, theme.accent)
        add_text(
            slide,
            fill_text(spec.table.row_heading, facts),
            left + 91440,
            top,
            label_w - 137160,
            row_h,
            font=theme.font_body,
            size=13,
            color=theme.cover_ink,
            bold=True,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        for index, (heading, _fact) in enumerate(columns):
            add_text(
                slide,
                heading,
                left + label_w + index * value_w,
                top,
                value_w - 91440,
                row_h,
                font=theme.font_body,
                size=13,
                color=theme.cover_ink,
                bold=True,
                align=PP_ALIGN.RIGHT,
                anchor=MSO_ANCHOR.MIDDLE,
            )
        for row_index, label in enumerate(labels):
            y = top + (row_index + 1) * row_h
            if row_index % 2 == 0:
                add_shape(slide, MSO_SHAPE.RECTANGLE, left, y, width, row_h, theme.card)
            add_text(
                slide,
                label,
                left + 91440,
                y,
                label_w - 137160,
                row_h,
                font=theme.font_body,
                size=14,
                color=theme.ink,
                anchor=MSO_ANCHOR.MIDDLE,
            )
            for index, (_heading, fact) in enumerate(columns):
                value = fact.values[row_index] if row_index < len(fact.values) else None
                shown = format_value(value, fact.format, fact.unit) if value is not None else ""
                add_text(
                    slide,
                    shown,
                    left + label_w + index * value_w,
                    y,
                    value_w - 91440,
                    row_h,
                    font=theme.font_display,
                    size=16,
                    color=theme.ink,
                    bold=True,
                    align=PP_ALIGN.RIGHT,
                    anchor=MSO_ANCHOR.MIDDLE,
                )
        paint_footer(slide, canvas, theme, number, _sources(spec, facts))
        return overflow


class SummaryLayout:
    layout = "summary"

    def paint(self, slide: Any, spec: Slide, facts: FactRegistry, theme: Theme, canvas: Canvas, number: int) -> bool:
        paint_background(slide, theme.paper)
        overflow = paint_action_title(slide, fill_text(spec.title, facts), theme, canvas)
        lines = [fill_text(bullet, facts) for bullet in spec.bullets]
        top = 1463040
        height = canvas.footer_top - top - 91440
        spans = card_spans(len(lines) or 1)
        for (start, count), line, index in zip(spans, lines, range(1, len(lines) + 1), strict=True):
            left, width = canvas.span(start, count)
            overflow = overflow or _takeaway(slide, line, theme, canvas, left, top, width, height, folio(index))
        overflow = overflow or too_dense(lines)
        paint_footer(slide, canvas, theme, number, _sources(spec, facts))
        return overflow


class SourcesLayout:
    layout = "appendix_sources"

    def paint(self, slide: Any, spec: Slide, facts: FactRegistry, theme: Theme, canvas: Canvas, number: int) -> bool:
        paint_background(slide, theme.paper)
        paint_action_title(slide, fill_text(spec.title, facts), theme, canvas)
        rows = [facts.get(fact_id) for fact_id in spec.sources if fact_id in facts]
        left, width = canvas.span(0, 12)
        top = 1600200
        available = canvas.footer_top - top - 91440
        if not rows:
            add_text(slide, "No sources", left, top, width, 365760, font=theme.font_body, size=16, color=theme.muted)
        else:
            row_h = min(731520, available // len(rows))
            size = 15 if row_h >= 640080 else 12
            label_width = int(width * 0.36)
            meta_left = left + 1005840 + label_width
            meta_width = max(1, width - (meta_left - left))
            for index, fact in enumerate(rows):
                y = top + index * row_h
                add_shape(slide, MSO_SHAPE.RECTANGLE, left, y + row_h - 9144, width, 9144, theme.rule)
                add_text(
                    slide,
                    f"{fact.id} ·",
                    left,
                    y,
                    914400,
                    row_h - 27432,
                    font=theme.font_display,
                    size=size,
                    color=theme.accent,
                    anchor=MSO_ANCHOR.MIDDLE,
                )
                add_text(
                    slide,
                    single_line(fact.label, canvas.inches(label_width), size),
                    left + 1005840,
                    y,
                    label_width,
                    row_h - 27432,
                    font=theme.font_body,
                    size=size,
                    color=theme.ink,
                    anchor=MSO_ANCHOR.MIDDLE,
                )
                meta = f"{fact.source.template_id}    {', '.join(fact.source.tables)}"
                add_text(
                    slide,
                    single_line(meta, canvas.inches(meta_width), max(11, size - 2)),
                    meta_left,
                    y,
                    meta_width,
                    row_h - 27432,
                    font=theme.font_body,
                    size=max(11, size - 2),
                    color=theme.muted,
                    anchor=MSO_ANCHOR.MIDDLE,
                )
        paint_footer(slide, canvas, theme, number, "")
        return False


def default_registry() -> LayoutRegistry:
    return LayoutRegistry(
        (
            TitleLayout(),
            SectionLayout(),
            AgendaLayout(),
            KpiLayout(),
            ChartLayout(),
            BulletsLayout(),
            TwoContentLayout(),
            ComparisonLayout(),
            TimelineLayout(),
            QuoteLayout(),
            TableLayout(),
            SummaryLayout(),
            SourcesLayout(),
        )
    )


def format_kpi(fact: Fact) -> str:
    if fact.value is None:
        raise GuardrailViolation(f"KPI needs a scalar fact: {fact.id}")
    return format_value(fact.value, fact.format, fact.unit)


def _series(spec: Slide, facts: FactRegistry) -> tuple[tuple[str, ...], tuple[tuple[str, tuple[float, ...]], ...], str]:
    assert spec.chart is not None
    plotted: list[tuple[str, tuple[float, ...]]] = []
    categories: tuple[str, ...] = ()
    fmt = "decimal2"
    for series in spec.chart.series:
        fact = facts.get(series.fact_id)
        if fact.kind != "series":
            raise GuardrailViolation(f"Chart needs a series fact: {fact.id}")
        categories = fact.categories
        fmt = fact.format
        plotted.append((fill_text(series.name, facts), fact.values))
    return categories, tuple(plotted), fmt


def _sources(spec: Slide, facts: FactRegistry) -> str:
    labels: list[str] = []
    for fact_id in spec.sources:
        if fact_id in facts:
            label = facts.get(fact_id).label
            if label not in labels:
                labels.append(label)
    return footer_sources(labels)


def _text_slide(
    slide: Any, spec: Slide, facts: FactRegistry, theme: Theme, canvas: Canvas, number: int, *, columns: int
) -> bool:
    paint_background(slide, theme.paper)
    overflow = paint_action_title(slide, fill_text(spec.title, facts), theme, canvas)
    lines = [fill_text(bullet, facts) for bullet in spec.bullets]
    if columns == 1:
        overflow = overflow or _rows(slide, lines, theme, canvas, numbered=False)
    else:
        midpoint = (len(lines) + 1) // 2
        top = 1463040
        height = canvas.footer_top - top - 91440
        panels = ((0, lines[:midpoint], "The path"), (6, lines[midpoint:], "The split"))
        for start, group, heading in panels:
            left, width = canvas.span(start, 6)
            overflow = overflow or _filled_panel(slide, heading, group, theme, canvas, left, top, width, height, "")
    overflow = overflow or too_dense(lines)
    paint_footer(slide, canvas, theme, number, _sources(spec, facts))
    return overflow


def _rows(slide: Any, lines: list[str], theme: Theme, canvas: Canvas, *, numbered: bool) -> bool:
    if not lines:
        return False
    left, width = canvas.span(0, 12)
    top = 1645920
    available = canvas.footer_top - top - 91440
    row_h = min(640080, available // len(lines))
    overflow = False
    for index, line in enumerate(lines):
        y = top + index * row_h
        wrapped = fit_paragraphs(
            [line], canvas.inches(width - 822960), canvas.inches(max(row_h - 45720, 1)), max_pt=18, min_pt=12
        )
        if numbered:
            add_text(
                slide,
                folio(index + 1),
                left,
                y,
                640080,
                row_h,
                font=theme.font_display,
                size=wrapped.font_pt,
                color=theme.accent,
                anchor=MSO_ANCHOR.MIDDLE,
            )
        else:
            add_shape(
                slide,
                MSO_SHAPE.OVAL,
                left,
                y + row_h // 2 - 27432,
                54864,
                54864,
                theme.accent,
            )
        add_text(
            slide,
            "\n".join(wrapped.lines),
            left + 822960,
            y,
            width - 822960,
            row_h,
            font=theme.font_body,
            size=wrapped.font_pt,
            color=theme.ink,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        overflow = overflow or wrapped.overflow
    return overflow


def _takeaway(
    slide: Any,
    line: str,
    theme: Theme,
    canvas: Canvas,
    left: int,
    top: int,
    width: int,
    height: int,
    index: str,
) -> bool:
    """One executive takeaway: a large index, a rule, and the sentence filling the card."""
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height, theme.card, radius=0.05)
    add_shape(slide, MSO_SHAPE.RECTANGLE, left, top, width, 54864, theme.accent)
    add_text(
        slide,
        index,
        left + 182880,
        top + 228600,
        width - 365760,
        1097280,
        font=theme.font_display,
        size=44,
        color=theme.accent,
        bold=True,
    )
    add_shape(slide, MSO_SHAPE.RECTANGLE, left + 182880, top + 1371600, 457200, 27432, theme.accent)
    fit = fit_paragraphs([line], canvas.inches(width - 365760), canvas.inches(height - 1828800), max_pt=22, min_pt=14)
    add_text(
        slide,
        "\n".join(fit.lines),
        left + 182880,
        top + 1554480,
        width - 365760,
        height - 1828800,
        font=theme.font_body,
        size=fit.font_pt,
        color=theme.ink,
    )
    return fit.overflow or too_dense([line])


def _filled_panel(
    slide: Any,
    heading: str,
    lines: list[str],
    theme: Theme,
    canvas: Canvas,
    left: int,
    top: int,
    width: int,
    height: int,
    index: str,
) -> bool:
    """A card that uses its full height: accent header, then one row per line."""
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height, theme.card, radius=0.05)
    header_h = min(640080, height // 4)
    add_shape(slide, MSO_SHAPE.RECTANGLE, left, top, width, header_h, theme.accent)
    if index:
        add_text(
            slide,
            index,
            left + 137160,
            top,
            640080,
            header_h,
            font=theme.font_display,
            size=22,
            color=theme.cover_ink,
            bold=True,
            anchor=MSO_ANCHOR.MIDDLE,
        )
    if heading:
        add_text(
            slide,
            heading,
            left + (777240 if index else 182880),
            top,
            width - (960120 if index else 365760),
            header_h,
            font=theme.font_display,
            size=18,
            color=theme.cover_ink,
            bold=theme.display_bold,
            anchor=MSO_ANCHOR.MIDDLE,
        )
    body_top = top + header_h
    body_h = height - header_h
    if not lines:
        return False
    row_h = body_h // len(lines)
    overflow = False
    for row, line in enumerate(lines):
        fit = fit_paragraphs(
            [line], canvas.inches(width - 411480), canvas.inches(max(row_h - 45720, 1)), max_pt=18, min_pt=12
        )
        y = body_top + row * row_h
        add_shape(slide, MSO_SHAPE.OVAL, left + 182880, y + row_h // 2 - 27432, 54864, 54864, theme.accent)
        add_text(
            slide,
            "\n".join(fit.lines),
            left + 320040,
            y,
            width - 457200,
            row_h,
            font=theme.font_body,
            size=fit.font_pt,
            color=theme.ink,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        overflow = overflow or fit.overflow
    return overflow or too_dense(lines)


def _kpi_card(
    slide: Any,
    kpi: Kpi,
    facts: FactRegistry,
    theme: Theme,
    canvas: Canvas,
    left: int,
    top: int,
    width: int,
    height: int,
) -> bool:
    fact = facts.get(kpi.fact_id)
    if fact.kind != "scalar" or fact.value is None:
        raise GuardrailViolation(f"KPI needs a scalar fact: {fact.id}")
    spark = facts.get(kpi.spark_fact_id) if kpi.spark_fact_id else None
    if spark is not None and spark.kind != "series":
        raise GuardrailViolation(f"Sparkline needs a series fact: {spark.id}")
    delta = facts.get(kpi.delta_fact_id) if kpi.delta_fact_id else None
    if delta is not None and (delta.kind != "scalar" or delta.value is None):
        raise GuardrailViolation(f"Delta needs a scalar fact: {delta.id if delta else kpi.delta_fact_id}")
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height, theme.card, radius=0.08)
    add_shape(slide, MSO_SHAPE.RECTANGLE, left, top, width, 36576, theme.accent)
    value = format_kpi(fact)
    value_fit = fit_paragraphs(
        [value], canvas.inches(width - 274320), 0.7, max_pt=36, min_pt=20, bold=True, leading=1.0
    )
    add_text(
        slide,
        value,
        left + 137160,
        top + 137160,
        width - 274320,
        731520,
        font=theme.font_display,
        size=value_fit.font_pt,
        color=theme.ink,
        bold=True,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    label = fill_text(kpi.label, facts)
    label_fit = fit_paragraphs([label], canvas.inches(width - 274320), 0.55, max_pt=14, min_pt=11)
    add_text(
        slide,
        "\n".join(label_fit.lines),
        left + 137160,
        top + 868680,
        width - 274320,
        457200,
        font=theme.font_body,
        size=label_fit.font_pt,
        color=theme.muted,
    )
    overflow = value_fit.overflow or label_fit.overflow
    note = format_kpi(delta) if delta is not None and delta.value is not None else ""
    if spark is not None:
        _sparkline(slide, spark.values, theme, left + 137160, top + height - 640080, width - 274320, 320040)
        if not note:
            note = direction_word(spark.values)
    if note:
        add_text(
            slide,
            note,
            left + 137160,
            top + height - 1005840,
            width - 274320,
            320040,
            font=theme.font_body,
            size=12,
            color=theme.accent,
            bold=True,
        )
    return overflow


def _sparkline(
    slide: Any, values: tuple[float, ...], theme: Theme, left: int, top: int, width: int, height: int
) -> None:
    """A column sparkline drawn as shapes, so a KPI slide still has no chart frame."""
    if len(values) < 2 or width < 2 or height < 2:
        return
    gap = max(1, width // (len(values) * 5))
    bar_w = max(1, (width - gap * (len(values) - 1)) // len(values))
    floor, peak = min(values), max(values)
    span = peak - floor or 1.0
    for index, value in enumerate(values):
        bar_h = max(27432, int((value - floor) / span * (height * 0.72) + height * 0.28))
        colour = theme.highlight if index == len(values) - 1 else theme.accent
        add_shape(
            slide,
            MSO_SHAPE.RECTANGLE,
            left + index * (bar_w + gap),
            top + height - bar_h,
            bar_w,
            bar_h,
            colour,
        )


def _callout(slide: Any, lines: list[str], theme: Theme, canvas: Canvas, top: int, height: int) -> bool:
    left, width = canvas.span(0, 12)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height, theme.card, radius=0.08)
    add_shape(slide, MSO_SHAPE.RECTANGLE, left, top, 54864, height, theme.accent)
    fit = fit_paragraphs(lines, canvas.inches(width - 457200), canvas.inches(height - 182880), max_pt=16, min_pt=12)
    add_text(
        slide,
        "\n".join(fit.lines),
        left + 274320,
        top + 91440,
        width - 411480,
        height - 182880,
        font=theme.font_body,
        size=fit.font_pt,
        color=theme.ink,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    return fit.overflow
