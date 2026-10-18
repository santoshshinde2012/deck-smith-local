"""Slide geometry and the small set of shapes every layout is built from.

The canvas is a 12-column grid. Painters ask for a column span instead of inventing margins.
"""

from __future__ import annotations

from dataclasses import dataclass
from typing import Any

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

from decksmith.adapters.local.design.fitting import fit_paragraphs, single_line
from decksmith.adapters.local.design.theme import Theme

EMU_PER_INCH = 914400


@dataclass(frozen=True)
class Canvas:
    """Widescreen (or template) slide size plus the shared grid."""

    width: int
    height: int
    columns: int = 12
    margin: int = 566928  # 0.62 in
    gutter: int = 164592  # 0.18 in

    @property
    def content_width(self) -> int:
        return self.width - 2 * self.margin

    @property
    def footer_top(self) -> int:
        return self.height - 420624  # 0.46 in

    def span(self, start: int, count: int) -> tuple[int, int]:
        """Left edge and width of ``count`` columns beginning at ``start`` (0-based)."""
        if count < 1 or start < 0 or start + count > self.columns:
            raise ValueError(f"column span {start}+{count} is outside 0..{self.columns}")
        usable = self.content_width
        gutters = (self.columns - 1) * self.gutter
        column = (usable - gutters) // self.columns
        leftover = usable - gutters - column * self.columns

        def edge(index: int) -> int:
            extra = min(index, leftover)
            gutters_before = min(index, self.columns - 1)
            return self.margin + index * column + extra + gutters_before * self.gutter

        return edge(start), edge(start + count) - edge(start)

    def inches(self, emu: int) -> float:
        return emu / EMU_PER_INCH


def card_spans(count: int) -> tuple[tuple[int, int], ...]:
    """Even column spans for ``count`` cards across the 12-column grid."""
    if count < 1:
        return ()
    base, extra = divmod(12, count)
    spans: list[tuple[int, int]] = []
    cursor = 0
    for index in range(count):
        width = base + (1 if index < extra else 0)
        spans.append((cursor, width))
        cursor += width
    return tuple(spans)


def rgb(hex_color: str) -> RGBColor:
    text = hex_color.removeprefix("#")
    red, green, blue = int(text[0:2], 16), int(text[2:4], 16), int(text[4:6], 16)
    return RGBColor(red, green, blue)  # type: ignore[no-untyped-call]


def paint_background(slide: Any, hex_color: str) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(hex_color)


def add_shape(
    slide: Any,
    kind: MSO_SHAPE,
    left: int,
    top: int,
    width: int,
    height: int,
    fill: str | None,
    *,
    line: str | None = None,
    line_pt: float = 1.25,
    radius: float | None = None,
) -> Any:
    shape = slide.shapes.add_shape(kind, Emu(left), Emu(top), Emu(width), Emu(height))
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_pt)
    if radius is not None:
        shape.adjustments[0] = radius
    return shape


def add_text(
    slide: Any,
    text: str,
    left: int,
    top: int,
    width: int,
    height: int,
    *,
    font: str,
    size: int,
    color: str,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    anchor: MSO_ANCHOR = MSO_ANCHOR.TOP,
) -> Any:
    box = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(max(width, 1)), Emu(max(height, 1)))
    frame = box.text_frame
    frame.word_wrap = True
    frame.auto_size = None
    frame.margin_left = Emu(0)
    frame.margin_right = Emu(0)
    frame.margin_top = Emu(0)
    frame.margin_bottom = Emu(0)
    frame.vertical_anchor = anchor
    lines = text.split("\n") if text else [""]
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        paragraph.line_spacing = 1.0
        paragraph.text = line
        if paragraph.runs:
            run = paragraph.runs[0]
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.name = font
            run.font.color.rgb = rgb(color)
    return box


def paint_action_title(slide: Any, title: str, theme: Theme, canvas: Canvas) -> bool:
    """Left-aligned sentence title with a short accent rule. Returns whether the title overflowed."""
    left, width = canvas.span(0, 12)
    fit = fit_paragraphs([title], canvas.inches(width), 0.92, max_pt=28, min_pt=18, leading=1.15)
    add_text(
        slide,
        "\n".join(fit.lines),
        left,
        347472,  # 0.38 in
        width,
        841248,  # 0.92 in
        font=theme.font_display,
        size=fit.font_pt,
        color=theme.ink,
        anchor=MSO_ANCHOR.BOTTOM,
    )
    rule_top = 1234440  # 1.35 in
    add_shape(slide, MSO_SHAPE.RECTANGLE, left, rule_top, 365760, 27432, theme.accent)  # 0.4 in by 0.03 in
    return fit.overflow


def paint_footer(
    slide: Any,
    canvas: Canvas,
    theme: Theme,
    number: int,
    sources: str,
    *,
    inverse: bool = False,
) -> None:
    """Slide number on every slide. Content slides also carry a one-line source line."""
    ink = theme.cover_muted if inverse else theme.muted
    strong = theme.cover_ink if inverse else theme.ink
    top = canvas.footer_top
    number_width = 640080  # 0.7 in
    if not inverse:
        left, width = canvas.span(0, 12)
        add_shape(slide, MSO_SHAPE.RECTANGLE, left, top, width, 9144, theme.rule)
    text_width = canvas.content_width - number_width - canvas.gutter
    if sources and not inverse:
        left, _ = canvas.span(0, 12)
        line = single_line(sources, canvas.inches(text_width), 12)
        add_text(
            slide,
            line,
            left,
            top + 64008,
            text_width,
            274320,
            font=theme.font_body,
            size=12,
            color=ink,
            anchor=MSO_ANCHOR.MIDDLE,
        )
    add_text(
        slide,
        folio(number),
        canvas.width - canvas.margin - number_width,
        top + 45720,
        number_width,
        292608,
        font=theme.font_display,
        size=13,
        color=strong,
        align=PP_ALIGN.RIGHT,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def folio(number: int) -> str:
    """Uppercase roman page mark.

    Arabic digits on a slide are reserved for figures that trace to a fact. The folio is chrome,
    so it is written as a roman numeral and never looks like a metric.
    """
    value = max(number, 1)
    glyphs = ((10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I"))
    parts: list[str] = []
    for amount, glyph in glyphs:
        while value >= amount:
            parts.append(glyph)
            value -= amount
    return "".join(parts)


def footer_sources(labels: list[str]) -> str:
    if not labels:
        return ""
    return "Sources   " + "   ·   ".join(labels)
