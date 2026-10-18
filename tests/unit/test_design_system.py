"""Themes, fitting, layout choice and on-slide geometry for the design system."""

from __future__ import annotations

import json
import shutil
import subprocess
from pathlib import Path
from typing import Any, get_args

import pytest
from pptx import Presentation
from pptx.util import Emu, Inches

from decksmith.adapters.local.design.canvas import Canvas, folio
from decksmith.adapters.local.design.charts import direct_labels, highlight_index, number_format
from decksmith.adapters.local.design.fitting import fit_paragraphs, single_line, too_dense, wrap
from decksmith.adapters.local.design.layouts import (
    ChartLayout,
    KpiLayout,
    LayoutRegistry,
    SourcesLayout,
    default_registry,
)
from decksmith.adapters.local.design.theme import THEMES, contrast_ratio, get_theme, relative_luminance
from decksmith.adapters.local.fake_llm import FakeLLM, _breakdown_title, _trend_title
from decksmith.adapters.local.pptx_renderer import PptxRenderer
from decksmith.bootstrap import Container
from decksmith.domain.composition import layout_for, narrative_roles
from decksmith.domain.deckspec import DeckSpec, Layout, Slide
from decksmith.domain.errors import GuardrailViolation
from decksmith.domain.facts import FactRegistry
from decksmith.domain.principal import Principal
from decksmith.ports.llm import LLMRequest, Message
from tests.conftest import REQUEST

_TEXT_PAIRS = (
    ("ink", "paper", 4.5),
    ("muted", "paper", 4.5),
    ("ink", "card", 4.5),
    ("muted", "card", 4.5),
    ("cover_ink", "cover", 4.5),
    ("cover_muted", "cover", 4.5),
)
_GRAPHIC_PAIRS = (("accent", "paper", 3.0), ("highlight", "paper", 3.0), ("accent", "card", 3.0))


def _on_slide(path: str) -> None:
    presentation = Presentation(path)
    for slide in presentation.slides:
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= presentation.slide_width
            assert shape.top + shape.height <= presentation.slide_height


def _planner(prompt: str, audience: str = "exec", feedback: str = "") -> dict[str, Any]:
    payload = {
        "request": prompt,
        "intake": {
            "deck_type": "retention_review",
            "audience": audience,
            "period": "2026-Q3",
            "template_id": "corp_default@1.0",
        },
        "templates": [],
        "feedback": feedback,
    }
    response = FakeLLM().complete(
        LLMRequest(
            node="planner",
            tier="small",
            run_id="r",
            messages=(Message(role="user", content=json.dumps(payload)),),
        )
    )
    parsed: dict[str, Any] = json.loads(response.text)
    return parsed


@pytest.mark.parametrize("theme_id", sorted(THEMES))
def test_theme_pairs_meet_wcag_contrast(theme_id: str) -> None:
    theme = get_theme(theme_id)
    for foreground, background, minimum in (*_TEXT_PAIRS, *_GRAPHIC_PAIRS):
        ratio = contrast_ratio(getattr(theme, foreground), getattr(theme, background))
        assert ratio >= minimum, f"{theme_id} {foreground}/{background} is {ratio:.2f}"
    for colour in theme.series:
        assert contrast_ratio(colour, theme.card) >= 3.0
        assert contrast_ratio(colour, theme.paper) >= 3.0


def test_unknown_theme_and_colour_fail_fast(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="unknown theme"):
        get_theme("no-such-theme")
    with pytest.raises(ValueError, match="unknown theme"):
        PptxRenderer(tmp_path, theme="no-such-theme")
    with pytest.raises(ValueError, match="#RRGGBB"):
        relative_luminance("blue")


def test_fitting_shrinks_wraps_and_flags_density() -> None:
    assert wrap("", 4, 14) == ()
    assert wrap("hello", 0, 14) == ("hello",)
    assert len(wrap("abcdefghij", 0.3, 18)) > 1
    empty = fit_paragraphs([""], 4, 1, max_pt=16, min_pt=12)
    assert empty.lines == () and empty.overflow is False
    fitted = fit_paragraphs(["Short line"], 6, 1, max_pt=20, min_pt=12)
    assert fitted.font_pt == 20 and fitted.overflow is False
    overflow = fit_paragraphs(["word " * 40], 1.0, 0.25, max_pt=20, min_pt=12)
    assert overflow.overflow is True and overflow.font_pt == 12
    assert single_line("", 2, 12) == ""
    assert single_line("Short", 4, 12) == "Short"
    assert single_line("one two three four five six seven", 0.5, 14).endswith("…")
    assert too_dense(["x" * 521])
    assert not too_dense(["short"])


def test_grid_spans_cover_the_content_box() -> None:
    canvas = Canvas(12_192_000, 6_858_000)
    left, width = canvas.span(0, 12)
    assert left == canvas.margin
    assert left + width == canvas.width - canvas.margin
    with pytest.raises(ValueError, match="column span"):
        canvas.span(10, 4)
    assert [folio(number) for number in (1, 4, 9, 14, 25)] == ["I", "IV", "IX", "XIV", "XXV"]


def test_highlight_and_label_rules() -> None:
    assert highlight_index("line", [1.0, 3.0, 2.0]) == 2
    assert highlight_index("bar", [1.0, 3.0, 3.0]) == 1
    assert highlight_index("column", []) == 0
    assert direct_labels(6, 1) and not direct_labels(7, 1) and not direct_labels(3, 2)
    assert number_format("pct1") == "0.0%"
    assert number_format("missing") == "General"


def test_roles_follow_audience_and_wording() -> None:
    assert narrative_roles("exec", "retention review") == ("cover", "scorecard", "trend", "breakdown", "sources")
    assert narrative_roles("board", "retention review")[1] == "divider"
    assert "agenda" in narrative_roles("exec", "Please add an agenda")
    assert "comparison" in narrative_roles("exec", "Compare retention versus the base")
    assert "timeline" in narrative_roles("exec", "Show the roadmap")
    assert "quote" in narrative_roles("exec", "Open with a quote")
    assert "split" in narrative_roles("exec", "Use a split view")
    assert layout_for("scorecard") == "kpi"
    with pytest.raises(ValueError, match="unknown narrative role"):
        layout_for("table")


def test_registry_covers_every_layout_and_can_be_replaced(
    tmp_path: Path, spec_dict: dict[str, object], facts: FactRegistry
) -> None:
    assert set(default_registry().layouts()) == set(get_args(Layout))

    class Marker:
        layout = "title"

        def paint(
            self, slide: object, spec: Slide, facts: FactRegistry, theme: object, canvas: Canvas, number: int
        ) -> bool:
            del spec, facts, theme, canvas, number
            box = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(2000000), Emu(400000))  # type: ignore[attr-defined]
            box.text_frame.text = "FROM-REGISTRY"
            return False

    rest = [painter for painter in default_registry()._painters.values() if painter.layout != "title"]  # type: ignore[attr-defined]
    painters = [Marker(), *rest]
    deck = PptxRenderer(tmp_path, registry=LayoutRegistry(painters)).render(
        DeckSpec.model_validate(spec_dict), facts, "reg"
    )
    presentation = Presentation(deck.path)
    texts = [shape.text_frame.text for shape in presentation.slides[0].shapes if shape.has_text_frame]
    assert "FROM-REGISTRY" in texts


@pytest.mark.parametrize("theme_id", sorted(THEMES))
def test_every_layout_stays_on_the_slide(tmp_path: Path, facts: FactRegistry, theme_id: str) -> None:
    other = facts.get("f3").model_copy(update={"id": "f4", "label": "Segment churn"})
    facts.add(other)
    spec = DeckSpec.model_validate(
        {
            "spec_version": "1.0",
            "template_id": "corp_default@1.0",
            "title": "Customer retention review",
            "audience": "exec",
            "period": "2026-Q3",
            "slides": [
                {"slide_id": "s01", "layout": "title", "title": "Customer retention review"},
                {
                    "slide_id": "s02",
                    "layout": "section",
                    "title": "The retention picture",
                    "subtitle": "For the exec audience",
                },
                {
                    "slide_id": "s03",
                    "layout": "agenda",
                    "title": "What this review covers",
                    "bullets": ["Figures", "Path", "Split"],
                },
                {
                    "slide_id": "s04",
                    "layout": "kpi",
                    "title": "Where retention stands",
                    "kpis": [{"label": "Churn rate", "fact_id": "f1"}, {"label": "Active", "fact_id": "f2"}],
                    "bullets": ["Churn reached {{fact:f1}}"],
                    "sources": ["f1", "f2"],
                },
                {
                    "slide_id": "s05",
                    "layout": "chart",
                    "title": "Monthly churn ended lower than it started",
                    "chart": {"type": "line", "series": [{"name": "Monthly churn", "fact_id": "f3"}]},
                    "sources": ["f3"],
                },
                {
                    "slide_id": "s06",
                    "layout": "chart",
                    "title": "Churn is highest in the leading segment",
                    "chart": {"type": "bar", "series": [{"name": "Segment churn", "fact_id": "f4"}]},
                    "sources": ["f4"],
                },
                {
                    "slide_id": "s07",
                    "layout": "chart",
                    "title": "Two series share one scale",
                    "chart": {
                        "type": "column",
                        "series": [
                            {"name": "Monthly churn", "fact_id": "f3"},
                            {"name": "Segment churn", "fact_id": "f4"},
                        ],
                    },
                    "sources": ["f3", "f4"],
                },
                {
                    "slide_id": "s08",
                    "layout": "bullets",
                    "title": "What the figures support",
                    "bullets": ["Read the path", "Read the split"],
                },
                {
                    "slide_id": "s09",
                    "layout": "two_content",
                    "title": "The story in two parts",
                    "bullets": ["Path", "Split", "Figures"],
                },
                {
                    "slide_id": "s10",
                    "layout": "comparison",
                    "title": "Two ways to read retention",
                    "columns": [
                        {"heading": "The rate", "bullets": ["Who left"]},
                        {"heading": "The base", "bullets": ["Who remains"]},
                    ],
                },
                {
                    "slide_id": "s11",
                    "layout": "timeline",
                    "title": "How the review is sequenced",
                    "milestones": [
                        {"label": "Headline", "detail": "The figures"},
                        {"label": "Trend", "detail": "The path"},
                    ],
                },
                {
                    "slide_id": "s12",
                    "layout": "quote",
                    "title": "The point of the review",
                    "quote": "A story about who stays.",
                },
                {
                    "slide_id": "s13",
                    "layout": "appendix_sources",
                    "title": "Sources",
                    "sources": ["f1", "f2", "f3", "f4"],
                },
            ],
        }
    )
    deck = PptxRenderer(tmp_path / theme_id, theme=theme_id).render(spec, facts, theme_id)
    assert deck.overflow_slides == ()
    _on_slide(deck.path)
    presentation = Presentation(deck.path)
    charts = [shape.chart for slide in presentation.slides for shape in slide.shapes if shape.has_chart]
    assert charts[0].value_axis.visible is False and charts[0].has_legend is False
    highlight = get_theme(theme_id).highlight.removeprefix("#")
    marker = charts[0].series[0].points[2].marker.format.fill.fore_color.rgb
    assert str(marker) == highlight
    assert charts[2].has_legend is True and charts[2].value_axis.visible is True
    texts = [shape.text_frame.text for shape in presentation.slides[-1].shapes if shape.has_text_frame]
    assert any(text.startswith("f1") for text in texts)
    assert any(text.startswith("f4") for text in texts)
    assert "SELECT" not in " ".join(texts)
    assert presentation.core_properties.category == theme_id


def test_template_sets_slide_size_and_shapes_stay_inside(
    tmp_path: Path, spec_dict: dict[str, object], facts: FactRegistry
) -> None:
    template = tmp_path / "wide.pptx"
    blank = Presentation()
    blank.slide_width, blank.slide_height = Inches(10), Inches(7.5)
    blank.save(str(template))
    deck = PptxRenderer(tmp_path / "out", template).render(DeckSpec.model_validate(spec_dict), facts, "sized")
    rendered = Presentation(deck.path)
    assert rendered.slide_width == Inches(10)
    _on_slide(deck.path)


def test_painters_reject_the_wrong_fact_kind(tmp_path: Path, facts: FactRegistry) -> None:
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Emu(12_192_000), Emu(6_858_000)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    width, height = presentation.slide_width, presentation.slide_height
    assert isinstance(width, int) and isinstance(height, int)
    canvas = Canvas(width, height)
    kpi = Slide.model_validate(
        {
            "slide_id": "s01",
            "layout": "kpi",
            "title": "Where retention stands",
            "kpis": [{"label": "Not a scalar", "fact_id": "f3"}],
        }
    )
    with pytest.raises(GuardrailViolation, match="scalar"):
        KpiLayout().paint(slide, kpi, facts, get_theme("meridian"), canvas, 1)
    chart = Slide.model_validate(
        {
            "slide_id": "s02",
            "layout": "chart",
            "title": "Needs a series",
            "chart": {"type": "line", "series": [{"name": "Churn", "fact_id": "f1"}]},
        }
    )
    with pytest.raises(GuardrailViolation, match="series"):
        ChartLayout().paint(slide, chart, facts, get_theme("meridian"), canvas, 2)
    empty = Slide.model_validate({"slide_id": "s03", "layout": "appendix_sources", "title": "Sources", "sources": []})
    blank = presentation.slides.add_slide(presentation.slide_layouts[6])
    assert SourcesLayout().paint(blank, empty, facts, get_theme("meridian"), canvas, 3) is False
    texts = [shape.text_frame.text for shape in blank.shapes if shape.has_text_frame]
    assert "No sources" in texts
    del tmp_path


def test_new_layouts_reject_incomplete_or_misplaced_content(spec_dict: dict[str, object]) -> None:
    spec = DeckSpec.model_validate(spec_dict)
    payload = spec.model_dump()
    payload["slides"][0]["quote"] = "A line"
    with pytest.raises(ValueError, match="quote"):
        DeckSpec.model_validate(payload)
    payload = spec.model_dump()
    payload["slides"].append(
        {"slide_id": "s06", "layout": "comparison", "title": "Compare", "columns": [{"heading": "Only", "bullets": []}]}
    )
    with pytest.raises(ValueError, match="two columns"):
        DeckSpec.model_validate(payload)
    payload = spec.model_dump()
    payload["slides"].append(
        {"slide_id": "s06", "layout": "timeline", "title": "Sequence", "milestones": [{"label": "One", "detail": ""}]}
    )
    with pytest.raises(ValueError, match="milestones"):
        DeckSpec.model_validate(payload)
    payload = spec.model_dump()
    payload["slides"].append({"slide_id": "s06", "layout": "quote", "title": "A line", "quote": ""})
    with pytest.raises(ValueError, match="quote"):
        DeckSpec.model_validate(payload)
    schema = DeckSpec.model_json_schema()
    assert "table" not in schema["$defs"]["Slide"]["properties"]["layout"]["enum"]
    assert "comparison" in schema["$defs"]["Slide"]["properties"]["layout"]["enum"]


def test_default_outline_ids_and_action_titles() -> None:
    outline = _planner(REQUEST)
    assert [slide["slide_id"] for slide in outline["slides"]] == ["s01", "s02", "s03", "s04", "s05"]
    edited = _planner(REQUEST, feedback="drop the segment slide")
    assert [slide["slide_id"] for slide in edited["slides"]] == ["s01", "s02", "s03", "s05"]
    board = _planner("Q2 2026 churn deck for the board", audience="board")
    assert [slide["layout"] for slide in board["slides"]] == [
        "title",
        "section",
        "kpi",
        "chart",
        "chart",
        "appendix_sources",
    ]
    facts = [
        {
            "id": "f1",
            "template_id": "churn_rate_quarter",
            "kind": "scalar",
            "values": [],
            "categories": [],
        },
        {
            "id": "f3",
            "template_id": "churn_trend_monthly",
            "kind": "series",
            "categories": ["2026-07", "2026-08", "2026-09"],
            "values": [0.024, 0.023, 0.021],
        },
        {
            "id": "f4",
            "template_id": "churn_by_segment",
            "kind": "series",
            "categories": ["Enterprise", "Mid-market", "SMB"],
            "values": [0.004, 0.016, 0.031],
        },
        {"id": "f2", "template_id": "nrr_quarter", "kind": "scalar", "values": [], "categories": []},
        {"id": "f5", "template_id": "active_subscriptions", "kind": "scalar", "values": [], "categories": []},
    ]
    payload = {"outline": outline, "facts": facts, "template_id": "corp_default@1.0", "fix_issues": []}
    written = json.loads(
        FakeLLM()
        .complete(
            LLMRequest(
                node="writer", tier="strong", run_id="r", messages=(Message(role="user", content=json.dumps(payload)),)
            )
        )
        .text
    )
    titles = {slide["slide_id"]: slide["title"] for slide in written["slides"]}
    assert titles["s03"] == "Monthly churn ended lower than it started"
    assert titles["s04"] == "Churn is highest in SMB"
    assert not any(character.isdigit() for title in titles.values() for character in title)
    assert _trend_title({"values": [1, 2]}) == "Monthly churn ended higher than it started"
    assert _trend_title({"values": [2, 2]}) == "Monthly churn ended where it started"
    assert _trend_title({"values": [1]}) == "The monthly path of churn"
    assert _breakdown_title({"categories": ["2026-07"], "values": [1]}, "fallback") == "Churn is uneven across segments"
    assert _breakdown_title({}, "fallback") == "fallback"


@pytest.mark.parametrize(
    ("prompt", "expected"),
    [
        ("Q2 2026 churn deck for the board", "section"),
        ("Q3 2026 retention review with an agenda", "agenda"),
        ("Compare retention versus the prior story for Q3 2026", "comparison"),
        ("Q3 2026 retention timeline", "timeline"),
        ("Q3 2026 retention review quote", "quote"),
        ("Q3 2026 retention split view", "two_content"),
    ],
)
def test_richer_layouts_survive_approval(prompt: str, expected: str, container: Container, bob: Principal) -> None:
    started = container.service.start(prompt, bob)
    assert started.outline is not None
    layouts = [slide["layout"] for slide in started.outline["slides"]]
    assert expected in layouts
    assert layouts[0] == "title" and layouts[-1] == "appendix_sources"
    done = container.service.decide(started.run_id, "approve", bob)
    assert done.status == "done"
    assert done.artifact is not None
    assert done.artifact["overflow_slides"] == []
    _on_slide(str(done.artifact["path"]))


def test_libreoffice_opens_the_deck(tmp_path: Path, spec_dict: dict[str, object], facts: FactRegistry) -> None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice is None:
        pytest.skip("LibreOffice is not installed")
    deck = PptxRenderer(tmp_path).render(DeckSpec.model_validate(spec_dict), facts, "office")
    completed = subprocess.run(
        [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(tmp_path), deck.path],
        check=False,
        capture_output=True,
        text=True,
        timeout=120,
    )
    assert completed.returncode == 0, completed.stderr
    pdfs = list(tmp_path.glob("*.pdf"))
    assert pdfs and pdfs[0].stat().st_size > 10_000
