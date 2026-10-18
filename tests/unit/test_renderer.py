from pathlib import Path

import pytest
from pptx import Presentation

from decksmith.adapters.local.design.charts import category_labels
from decksmith.adapters.local.design.text import fill_text
from decksmith.adapters.local.pptx_renderer import PptxRenderer
from decksmith.domain.deckspec import DeckSpec
from decksmith.domain.errors import GuardrailViolation
from decksmith.domain.facts import FactRegistry


def test_renders_native_objects(renderer: PptxRenderer, spec_dict: dict[str, object], facts: FactRegistry) -> None:
    deck = renderer.render(DeckSpec.model_validate(spec_dict), facts, "run1")
    prs = Presentation(deck.path)
    assert deck.slide_count == len(prs.slides) == 5
    texts = [sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame]
    assert any("2.4%" in t for t in texts) and any("6,172" in t for t in texts)
    assert not any("{{fact" in t for t in texts)
    charts = [sh.chart for s in prs.slides for sh in s.shapes if sh.has_chart]
    assert len(charts) == 1
    assert list(charts[0].plots[0].categories) == category_labels(["2026-07", "2026-08", "2026-09"])
    visible = "\n".join(text for text in texts)
    assert "queried" not in visible and "2026-" not in visible
    assert prs.core_properties.subject == "run_id=run1"
    assert "f3" in prs.slides[2].notes_slide.notes_text_frame.text  # provenance in speaker notes


def test_content_hash_is_stable(tmp_path: Path, spec_dict: dict[str, object], facts: FactRegistry) -> None:
    spec = DeckSpec.model_validate(spec_dict)
    a = PptxRenderer(tmp_path / "a").render(spec, facts, "same")
    b = PptxRenderer(tmp_path / "b").render(spec, facts, "same")
    assert a.content_hash == b.content_hash


def test_overflow_flagged(renderer: PptxRenderer, spec_dict: dict[str, object], facts: FactRegistry) -> None:
    spec_dict["slides"][3]["bullets"] = ["word " * 30] * 6  # type: ignore[index]
    deck = renderer.render(DeckSpec.model_validate(spec_dict), facts, "run2")
    assert deck.overflow_slides == ("s04",)


def test_placeholders_expand_in_subtitles_labels_legends_and_notes(
    renderer: PptxRenderer, spec_dict: dict[str, object], facts: FactRegistry
) -> None:
    spec = DeckSpec.model_validate(spec_dict)
    spec.title = "Review {{fact:f1}}"
    spec.slides[0].subtitle = "Churn {{fact:f1}}"
    spec.slides[1].kpis[0].label = "Churn {{fact:f1}}"
    spec.slides[3].bullets = ["Churn {{fact:f1}}"]
    spec.slides[3].notes = "Active {{fact:f2}}"
    assert spec.slides[2].chart is not None
    spec.slides[2].chart.series[0].name = "Churn {{fact:f1}}"
    deck = renderer.render(spec, facts, "placeholders")
    prs = Presentation(deck.path)
    texts = [shape.text_frame.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame]
    assert "Churn 2.4%" in texts
    assert any("Churn 2.4%" in text for text in texts)
    chart = next(shape.chart for shape in prs.slides[2].shapes if shape.has_chart)
    assert chart.series[0].name == "Churn 2.4%"
    notes = prs.slides[3].notes_slide.notes_text_frame.text
    assert "Active 6,172" in notes
    assert "f1 · Churn rate" in notes and "f2 · Active subscriptions" in notes
    assert prs.core_properties.title == "Review 2.4%"
    assert all("{{fact:" not in text for text in [*texts, notes])


def test_renderer_refuses_series_in_text_and_missing_source_appendix(
    renderer: PptxRenderer, spec_dict: dict[str, object], facts: FactRegistry
) -> None:
    spec = DeckSpec.model_validate(spec_dict)
    spec.slides[3].bullets = ["Churn {{fact:f3}}"]
    with pytest.raises(GuardrailViolation, match="placeholder_needs_scalar"):
        renderer.render(spec, facts, "bad-series")
    with pytest.raises(GuardrailViolation, match="needs a scalar"):
        fill_text("Churn {{fact:f3}}", facts)
    spec.slides[3].bullets = []
    spec.slides.pop()
    with pytest.raises(GuardrailViolation, match="missing_sources_appendix"):
        renderer.render(spec, facts, "no-sources")


def test_query_details_are_in_notes_with_query_time_not_freshness(
    renderer: PptxRenderer, spec_dict: dict[str, object], facts: FactRegistry
) -> None:
    source = facts.get("f1").source.model_copy(
        update={
            "sql": "SELECT 0.0236 WHERE month >= :start_month",
            "params": {"start_month": "2026-07-01"},
            "template_sha256": "abc123",
        }
    )
    detailed = FactRegistry(
        [fact.model_copy(update={"source": source}) if fact.id == "f1" else fact for fact in facts.all()]
    )
    deck = renderer.render(DeckSpec.model_validate(spec_dict), detailed, "query-details")
    prs = Presentation(deck.path)
    notes = prs.slides[1].notes_slide.notes_text_frame.text
    assert "SELECT 0.0236 WHERE month >= :start_month" in notes
    assert '"start_month": "2026-07-01"' in notes
    assert "template SHA-256: abc123" in notes
    assert "queried 2026-10-01" in notes
    source_text = "\n".join(shape.text_frame.text for shape in prs.slides[-1].shapes if shape.has_text_frame)
    assert "SELECT" not in source_text
    assert "queried" not in source_text


def test_section_subtitle_is_rendered_and_cited(
    renderer: PptxRenderer, spec_dict: dict[str, object], facts: FactRegistry
) -> None:
    spec = DeckSpec.model_validate(spec_dict)
    spec.slides[0].layout = "section"
    spec.slides[0].subtitle = "Churn {{fact:f1}}"
    deck = renderer.render(spec, facts, "section-subtitle")
    prs = Presentation(deck.path)
    assert any(shape.text_frame.text == "Churn 2.4%" for shape in prs.slides[0].shapes if shape.has_text_frame)
    assert "f1 · Churn rate" in prs.slides[0].notes_slide.notes_text_frame.text
